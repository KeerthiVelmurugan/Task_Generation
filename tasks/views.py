from django.shortcuts import render

# Create your views here.
import random
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from docx import Document
from .models import Concept
from rest_framework import generics
from .models import Task
from .serializers import TaskSerializer,ConceptSerializer
from rest_framework.exceptions import NotFound



# GET, POST, PUT (Edit), DELETE for Concept
class ConceptListCreateView(APIView):
    # GET method to list all concepts
    def get(self, request):
        concepts = Concept.objects.all()
        serializer = ConceptSerializer(concepts, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)

    # POST method to create a new concept
    def post(self, request):
        serializer = ConceptSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

class ConceptDetailView(APIView):
    # GET method to retrieve a specific concept by ID
    def get(self, request, pk):
        try:
            concept = Concept.objects.get(pk=pk)
        except Concept.DoesNotExist:
            return Response({'error': 'Concept not found.'}, status=status.HTTP_404_NOT_FOUND)
        serializer = ConceptSerializer(concept)
        return Response(serializer.data)

    # PUT method to edit a specific concept by ID
    def put(self, request, pk):
        try:
            concept = Concept.objects.get(pk=pk)
        except Concept.DoesNotExist:
            return Response({'error': 'Concept not found.'}, status=status.HTTP_404_NOT_FOUND)
        
        serializer = ConceptSerializer(concept, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    # DELETE method to remove a specific concept by ID
    def delete(self, request, pk):
        try:
            concept = Concept.objects.get(pk=pk)
        except Concept.DoesNotExist:
            return Response({'error': 'Concept not found.'}, status=status.HTTP_404_NOT_FOUND)

        concept.delete()
        return Response({'message': 'Concept deleted successfully.'}, status=status.HTTP_204_NO_CONTENT)



class GenerateTasksDocumentView(APIView):
    def post(self, request):
        concept_id = request.data.get('concept_id')
        if not concept_id:
            return Response({'error': 'Concept ID is required.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            concept = Concept.objects.get(id=concept_id)
        except Concept.DoesNotExist:
            return Response({'error': 'Concept not found.'}, status=status.HTTP_404_NOT_FOUND)

        tasks = list(concept.tasks.all())
        if len(tasks) < 5:
            return Response({'error': 'Not enough tasks under this concept.'}, status=status.HTTP_400_BAD_REQUEST)

        selected_tasks = random.sample(tasks, 5)

        # Create Word document
        document = Document()
        document.add_heading(f'Tasks for Concept: {concept.name}', level=1)

        for idx, task in enumerate(selected_tasks, start=1):
            document.add_paragraph(f"{idx}. {task.description}")

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = f'attachment; filename=Tasks_{concept.name}.docx'
        document.save(response)

        return response
    


def export_random_tasks(request, concept_id):
    try:
        concept = Concept.objects.get(id=concept_id)
    except Concept.DoesNotExist:
        return HttpResponse("Concept not found.", status=404)

    tasks = list(concept.tasks.all())

    if len(tasks) < 5:
        return HttpResponse("Not enough tasks to select 5.", status=400)

    random_tasks = random.sample(tasks, 5)

    # Create Word document
    document = Document()
    document.add_heading(f'Tasks for Concept: {concept.name}', level=1)

    for idx, task in enumerate(random_tasks, start=1):
        document.add_paragraph(f"{idx}. {task.description}")

    # Prepare HTTP response
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = f'attachment; filename={concept.name}_tasks.docx'
    document.save(response)

    return response






class TasksByConceptView(generics.ListAPIView):
    serializer_class = TaskSerializer
    
    def get_queryset(self):
        concept_id = self.kwargs['concept_id']
        tasks = Task.objects.filter(concept_id=concept_id)
        if not tasks.exists():
            raise NotFound(detail="Concept not found or no tasks available")
        return tasks

    def list(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        task_count = queryset.count()
        serializer = self.get_serializer(queryset, many=True)
        
        # Fetch the concept object to get the concept name
        concept = Concept.objects.get(id=kwargs['concept_id'])
        
        # Modify each task data to include the concept name
        for task_data in serializer.data:
            task_data['concept_name'] = concept.name  # Add concept name instead of concept ID
            del task_data['concept']  # Optionally, remove concept ID

        # Create the final response data
        response_data = {
            'concept_id': kwargs['concept_id'],
            'concept_name': concept.name,  # Add concept name to the response
            'task_count': task_count,
            'tasks': serializer.data
        }

        return Response(response_data)
    
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
import random
from .models import Task
from .serializers import TaskSerializer

class RandomTasksFromAllConceptsView(APIView):
    def get(self, request):
        # Get the 'count' from query params, default to 5 if not provided
        task_count = int(request.query_params.get('count', 5))

        tasks = list(Task.objects.all())

        if len(tasks) < task_count:
            return Response({'error': f'Not enough tasks in the database. Available tasks: {len(tasks)}.'},
                            status=status.HTTP_400_BAD_REQUEST)

        random_tasks = random.sample(tasks, task_count)

        serializer = TaskSerializer(random_tasks, many=True)
        return Response(serializer.data, status=status.HTTP_200_OK)


class RandomTasksExportView(APIView):
    def get(self, request):
        task_count = int(request.GET.get('count', 5))  # Get task count from query param, default to 5
        tasks = Task.objects.all()  # Get all tasks

        if len(tasks) < task_count:
            return Response({"error": "Not enough tasks available."}, status=status.HTTP_400_BAD_REQUEST)

        # Randomly select the tasks
        random_tasks = random.sample(list(tasks), task_count)

        # If export is requested
        if 'export' in request.GET:
            # Create Word document
            document = Document()
            document.add_heading(f'Random Tasks (Total {task_count})', level=1)

            for idx, task in enumerate(random_tasks, start=1):
                document.add_paragraph(f"{idx}. {task.description}")

            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = f'attachment; filename=Random_Tasks_{task_count}.docx'
            document.save(response)
            return response

        # If export is not requested, just return the task data
        task_data = [{'task_id': task.id, 'description': task.description} for task in random_tasks]
        return Response(task_data)

# views.py
import json
from django.views import View
from django.http import JsonResponse
from django.core.mail import EmailMessage
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator
from io import BytesIO
from docx import Document
from .models import Task

@method_decorator(csrf_exempt, name='dispatch')
class SendTasksToEmailAPIView(View):
    def post(self, request):
        try:
            data = json.loads(request.body)
            email = data.get('email')
            count = int(data.get('count', 5))

            if not email:
                return JsonResponse({'error': 'Email is required.'}, status=400)

            tasks = Task.objects.order_by('?')[:count]

            doc = Document()
            for idx, task in enumerate(tasks, start=1):
                doc.add_paragraph(f"{idx}. {task.description}")

            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            mail = EmailMessage(
                subject='Your Random Tasks',
                body='Please find the attached task file.',
                to=[email],
            )
            mail.attach('tasks.docx', buffer.read(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            mail.send()

            return JsonResponse({'message': 'Tasks sent successfully to your email.'})
        
        except json.JSONDecodeError:
            return JsonResponse({'error': 'Invalid JSON'}, status=400)
        except Exception as e:
            return JsonResponse({'error': str(e)}, status=500)




